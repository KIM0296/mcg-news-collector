"""기존 '국내외단신DB' 엑셀과 같은 모양으로 저장한다.

A열·1행은 비우고, 2행에 머리글: 구분 | 보도일자 | 매체명 | 제 목 | 핵심 키워드 | 관련국가 | 원문링크 | 원문 요약
(앞의 7개 열은 기존 DB와 같다. '원문 요약'은 근거 확인용으로 추가한 열)
"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

HEADERS = ["구분", "보도일자", "매체명", "제 목", "핵심 키워드", "관련국가", "원문링크", "원문 요약"]
WIDTHS = [14, 12, 20, 60, 40, 12, 45, 60]
CATEGORY_ORDER = {"국내기사": 0, "사설/칼럼/기고": 1, "세계기사": 2, "해외단신": 3, "WNN": 4, "원자력계소식": 5}


def save(items, path, sheet_title):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]
    thin = Side(style="thin", color="BFBFBF")
    head_fill = PatternFill("solid", fgColor="841938")
    for j, (h, w) in enumerate(zip(HEADERS, WIDTHS), start=2):
        c = ws.cell(row=2, column=j, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
        ws.column_dimensions[c.column_letter].width = w
    ws.column_dimensions["A"].width = 3

    rows = sorted(items, key=lambda i: (i.date, CATEGORY_ORDER.get(i.category, 9)))
    for r, it in enumerate(rows, start=3):
        vals = [it.category, it.date, it.media, it.title, it.keywords, it.country, it.url, it.summary]
        for j, v in enumerate(vals, start=2):
            c = ws.cell(row=r, column=j, value=v)
            c.alignment = Alignment(vertical="top", wrap_text=j in (5, 6, 9))
            c.border = Border(bottom=thin)
        link = ws.cell(row=r, column=8)
        if it.url:
            link.hyperlink = it.url
            link.font = Font(color="0563C1", underline="single")
    ws.freeze_panes = "B3"
    ws.auto_filter.ref = f"B2:I{max(2, len(rows) + 2)}"
    wb.save(path)
    return path
